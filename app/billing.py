"""PDF bills + tax table."""

from __future__ import annotations

import io
import re
from collections import defaultdict
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from xml.sax.saxutils import escape

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.platypus import Paragraph, Table, TableStyle
from reportlab.pdfgen import canvas

from . import config, storage

_MONEY = Decimal("0.01")
_TAX_RATE = Decimal(str(config.CT_SALES_TAX_RATE))
_THANKS = (
    "Thank you for your business. I look forward to continuing my work with you. "
    "Let me know if you have any questions or anything else I can help you with. "
    "Please make the checks payable to Nick's Lawn Care LLC."
)
_MOW = re.compile(r"(?i)^mow(?:ing)?(?:\s+(\d{1,2}))?$")
_HEDGE = re.compile(r"(?i)^hedg(?:e|ing)(?:\s+(\d{1,2}))?$")
_NOTE_DAY = re.compile(r"(?i)^(\d{1,2})h?$")
_PRIOR_NOTE = re.compile(r"(?i)^prior(?P<taxed>-taxed)?:(?P<month>.*)$")
_EMAIL = re.compile(r"[A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,}", re.IGNORECASE)


def _money(amount) -> Decimal:
    return Decimal(str(amount or 0)).quantize(_MONEY, rounding=ROUND_HALF_UP)


def _money_label(amount: Decimal) -> str:
    if amount == amount.to_integral():
        return f"${int(amount)}"
    return f"${amount:.2f}"


def _ordinal(day: int) -> str:
    if 10 <= day % 100 <= 20:
        suffix = "th"
    else:
        suffix = {1: "st", 2: "nd", 3: "rd"}.get(day % 10, "th")
    return f"{day}{suffix}"


def _first_name(name: str) -> str:
    """Given name only. Roster rows are stored LAST, First."""
    raw = " ".join(str(name or "").split())
    if not raw:
        return "Customer"
    if "," in raw:
        given = raw.split(",", 1)[1].strip()
    else:
        parts = raw.split()
        if len(parts) >= 2 and parts[0].isupper():
            given = parts[-1]
        else:
            given = raw
    parts = [part for part in given.split() if part not in {"&", "and"}]
    word = parts[0] if parts else ""
    if not word:
        return "Customer"
    if word.isupper():
        word = word.title()
    return word


def _invoice_when(month: str) -> tuple[str, str, int]:
    """Letter date, work-month name, and work-month number.

    A YYYY-MM month is billed on the first of the next month. Any other label,
    such as a test month, uses today's date and that label as the period.
    """
    raw = str(month or "").strip()
    parts = raw.split("-", 1)
    if len(parts) == 2 and parts[0].isdigit() and parts[1].isdigit():
        year, mon = int(parts[0]), int(parts[1])
        if 1 <= mon <= 12:
            work = date(year, mon, 1)
            if mon == 12:
                issued = date(year + 1, 1, 1)
            else:
                issued = date(year, mon + 1, 1)
            letter = f"{issued.strftime('%B')} {_ordinal(issued.day)}, {issued.year}"
            return letter, work.strftime("%B"), mon
    today = date.today()
    letter = f"{today.strftime('%B')} {_ordinal(today.day)}, {today.year}"
    label = raw.replace("_", " ") or "this period"
    return letter, label, 0


def invoice_sentence(month: str, custom: str | None = None) -> str:
    """The line under Dear. A saved sentence replaces the month line."""
    text = _letter_text(custom)
    if text:
        return text
    _letter, work_month, _month_num = _invoice_when(month)
    return f"Below is the invoice for any work done in {work_month}."


def default_greeting(client_name: str) -> str:
    return f"Dear {_first_name(client_name)},"


def default_closing() -> str:
    return _THANKS


def default_signoff() -> str:
    return f"With regards,\n{config.INVOICE_SIGNOFF}"


def _letter_text(value: str | None) -> str:
    return "\n".join(str(value or "").replace("\r\n", "\n").replace("\r", "\n").splitlines()).strip()


def letter_or_default(custom: str | None, default: str) -> str:
    text = _letter_text(custom)
    return text or default


def stored_letter(custom: str | None, default: str) -> str | None:
    """None when the editor still has the usual wording, so a later default can change."""
    text = _letter_text(custom)
    if not text or text == _letter_text(default):
        return None
    return text


def _service(line: dict) -> tuple[str, int | None]:
    """Job name and day-of-month for one stored work row."""
    desc = str(line.get("description") or "").strip()
    note = str(line.get("day_or_note") or "").strip()
    day = None
    noted = _NOTE_DAY.fullmatch(note)
    if noted:
        day = int(noted.group(1))
    mow = _MOW.fullmatch(desc)
    if mow:
        return "Mowing", day if day is not None else (int(mow.group(1)) if mow.group(1) else None)
    hedge = _HEDGE.fullmatch(desc)
    if hedge:
        return "Hedging", day if day is not None else (int(hedge.group(1)) if hedge.group(1) else None)
    return desc, day


def line_role(line: dict) -> tuple[str, str]:
    """visit, discount, prior, or prior_taxed, plus the prior month label."""
    note = str(line.get("day_or_note") or "").strip()
    desc = str(line.get("description") or "").strip()
    if note.lower() == "discount" or desc.lower().startswith("discount"):
        return "discount", ""
    prior = _PRIOR_NOTE.match(note)
    if prior:
        role = "prior_taxed" if prior.group("taxed") else "prior"
        return role, (prior.group("month") or "").strip()
    kind = str(line.get("kind") or "").strip()
    if kind in {"discount", "prior", "prior_taxed"}:
        return kind, str(line.get("prior_month") or "").strip()
    return "visit", ""


def contact_email(row: dict) -> str:
    """Addresses to email. Notes count only when the email field is empty."""
    email = (row.get("email") or "").strip()
    if "@" in email:
        return email
    return ", ".join(_EMAIL.findall(row.get("billing_notes") or ""))


def delivery_channel(row: dict) -> str:
    """A saved choice wins. Otherwise Mail is paper, then email, then SMS, then paper."""
    choice = str(row.get("delivery") or "").strip().lower()
    if choice in {"email", "sms", "mail"}:
        return choice
    if row.get("prefer_mail"):
        return "mail"
    if "@" in contact_email(row):
        return "email"
    if (row.get("phone") or "").strip():
        return "sms"
    return "mail"


def face_bill(row: dict) -> dict:
    """Use this month's name, email, and delivery when the bill has its own."""
    shown = dict(row)
    name = str(shown.get("display_name") or "").strip()
    if name:
        shown["client_name"] = name
    if shown.get("bill_email") is not None:
        shown["email"] = shown.get("bill_email") or ""
    choice = str(shown.get("bill_delivery") or "").strip().lower()
    if choice in {"email", "sms", "mail"}:
        shown["delivery"] = choice
    return shown


def compile_invoice_lines(lines: list[dict], month: str) -> list[dict]:
    """Group visits into Date / Description / Price rows, then sales tax and total.

    Every mowing day in the month shares one row. Every hedge day shares one row.
    A written job keeps its name. A discount reduces the taxable subtotal and has
    no date. A previous-month total is taxed with this month unless it is marked
    as already including sales tax, in which case it is added after the tax.
    """
    _letter, _work_name, month_num = _invoice_when(month)
    groups: list[dict] = []
    index: dict[str, dict] = {}
    for line in lines:
        role, _prior_month = line_role(line)
        label, day = _service(line)
        if not label:
            continue
        if label.lower() in {"sales tax", "total"}:
            continue
        key = f"{role}:{label.casefold()}"
        row = index.get(key)
        if row is None:
            row = {"description": label, "days": [], "amount": Decimal("0"), "role": role}
            index[key] = row
            groups.append(row)
        if role == "visit" and day is not None and 1 <= day <= 31 and day not in row["days"]:
            row["days"].append(day)
        row["amount"] += _money(line.get("amount"))
    table = []
    taxable = Decimal("0")
    after_tax = Decimal("0")
    taxed_rows = []
    for row in groups:
        amount = row["amount"].quantize(_MONEY, rounding=ROUND_HALF_UP)
        if row["role"] == "discount":
            amount = -abs(amount)
        days = sorted(row["days"])
        if row["role"] == "visit" and days:
            if month_num:
                dates = ", ".join(f"{month_num}/{day}" for day in days)
            else:
                dates = ", ".join(str(day) for day in days)
        else:
            dates = ""
        printed = {"date": dates, "description": row["description"], "amount": amount}
        if row["role"] == "prior_taxed":
            taxed_rows.append(printed)
            after_tax += amount
        else:
            table.append(printed)
            taxable += amount
    tax = (taxable * _TAX_RATE).quantize(_MONEY, rounding=ROUND_HALF_UP)
    total = (taxable + tax + after_tax).quantize(_MONEY, rounding=ROUND_HALF_UP)
    table.append({"date": "", "description": "Sales tax", "amount": tax})
    table.extend(taxed_rows)
    table.append({"date": "", "description": "Total", "amount": total})
    return table


def _wrap(text: str, font: str, size: float, width: float) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        trial = word if not current else f"{current} {word}"
        if stringWidth(trial, font, size) <= width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines or [""]


def build_pdf_bytes(*, company: str, client_name: str, address: str, month: str,
                    lines: list[dict], email: str = "", intro: str | None = None,
                    greeting: str | None = None, closing: str | None = None,
                    signoff: str | None = None) -> bytes:
    del company, address, email
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    page_w, page_h = letter
    left = 0.75 * inch
    right = page_w - 0.75 * inch
    width = right - left
    rows = compile_invoice_lines(lines, month)
    letter_date, _work_month, _month_num = _invoice_when(month)

    cell = ParagraphStyle(
        "invCell", fontName="Times-Roman", fontSize=11, leading=14,
        alignment=TA_LEFT, textColor=colors.black,
    )
    cell_center = ParagraphStyle("invCenter", parent=cell, alignment=TA_CENTER)
    cell_right = ParagraphStyle("invRight", parent=cell, alignment=TA_RIGHT)
    data = [[
        Paragraph("Date", cell_center),
        Paragraph("Description", cell_center),
        Paragraph("Price", cell_center),
    ]]
    for row in rows:
        data.append([
            Paragraph(escape(row["date"]), cell),
            Paragraph(escape(row["description"]), cell),
            Paragraph(escape(_money_label(row["amount"])), cell_right),
        ])
    table = Table(data, colWidths=[width * 0.34, width * 0.46, width * 0.20])
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Times-Roman"),
        ("FONTSIZE", (0, 0), (-1, -1), 11),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.8, colors.black),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("BACKGROUND", (0, 0), (-1, -1), colors.white),
    ]))
    _tw, table_h = table.wrap(width, page_h)

    def footer() -> None:
        c.setFillColor(colors.black)
        c.setFont("Times-Italic", 11)
        y = 0.95 * inch
        for line in (config.SERVICE_LINE, config.TAX_REGISTRATION):
            c.drawCentredString(page_w / 2, y, line)
            y -= 14
        y -= 8
        c.drawCentredString(page_w / 2, y, config.INSURED_LINE)

    def letterhead(y: float) -> float:
        c.setFillColor(colors.black)
        c.setFont("Times-Roman", 12)
        for line in (
            config.OWNER_NAME,
            config.COMPANY_NAME,
            config.COMPANY_ADDRESS,
            f"{config.COMPANY_PHONE}  {config.COMPANY_EMAIL}",
        ):
            c.drawCentredString(page_w / 2, y, line)
            y -= 15
        return y - 22

    y = letterhead(page_h - 0.7 * inch)
    c.setFont("Times-Roman", 12)
    c.drawString(left, y, letter_date)
    y -= 22

    def draw_text(text: str, y_pos: float, *, center: bool, width_limit: float) -> float:
        c.setFont("Times-Roman", 12)
        paragraphs = str(text).splitlines() or [""]
        for para in paragraphs:
            wrapped = _wrap(para, "Times-Roman", 12, width_limit) if para.strip() else [""]
            for line in wrapped:
                if center and line:
                    c.drawCentredString(page_w / 2, y_pos, line)
                else:
                    c.drawString(left, y_pos, line)
                y_pos -= 16
        return y_pos

    y = draw_text(letter_or_default(greeting, default_greeting(client_name)), y, center=False, width_limit=width)
    y -= 4
    y = draw_text(invoice_sentence(month, intro), y, center=False, width_limit=width)
    y -= 6
    table.drawOn(c, left, y - table_h)
    y = y - table_h - 28
    y = draw_text(letter_or_default(closing, default_closing()), y, center=True, width_limit=width * 0.92)
    y -= 8
    draw_text(letter_or_default(signoff, default_signoff()), y, center=False, width_limit=width)
    footer()
    c.showPage()
    c.save()
    return buf.getvalue()


def _reprice_visits(conn, lines: list[dict]) -> list[dict]:
    """Mowing and hedging use the client's current prices. Other lines stay."""
    from . import db

    for line in lines:
        role, _prior = line_role(line)
        if role != "visit":
            continue
        label, _day = _service(line)
        price = None
        if label == "Mowing":
            price = line.get("mow_price")
        elif label == "Hedging":
            price = line.get("hedge_price")
        if price is None:
            continue
        amount = float(price)
        if float(line.get("amount") or 0) != amount:
            db.update_work_item(conn, int(line["id"]), description=line.get("description") or label, amount=amount)
            line["amount"] = amount
    return lines


def _retarget_month_clients(conn, month: str) -> None:
    """Point this month's work at the set-up client when the sheet spelling differs."""
    from . import db
    from .knowledge import client_is_set_up, filing_client

    roster = [row for row in db.list_clients(conn) if row.get("on_roster") is not False]
    names = {}
    for row in db.work_for_month(conn, month):
        names[int(row["client_id"])] = row["client_name"]
    for cid, name in names.items():
        hit = filing_client(name, roster)
        if not hit or int(hit["id"]) == cid:
            continue
        db.reassign_month_client(conn, month, cid, int(hit["id"]))
        current = db.get_client(conn, cid)
        if not current or client_is_set_up(current):
            continue
        leftover = conn.execute(
            "SELECT 1 FROM work_items WHERE client_id = %s LIMIT 1",
            (cid,),
        ).fetchone()
        if not leftover:
            db.delete_client(conn, cid)


def generate_month_bills(conn, month: str) -> list[dict]:
    """Create PDFs for each client with work in month; upload to S3; save bill rows."""
    from . import db

    _retarget_month_clients(conn, month)
    rows = db.work_for_month(conn, month)
    by_client: dict[int, list] = defaultdict(list)
    meta: dict[int, dict] = {}
    for r in rows:
        cid = int(r["client_id"])
        by_client[cid].append(dict(r))
        meta[cid] = r
    out = []
    for cid, lines in by_client.items():
        m = meta[cid]
        _reprice_visits(conn, lines)
        letter = db.bill_letter(conn, month, cid)
        shown = str(letter.get("display_name") or "").strip() or m["client_name"]
        pdf = build_pdf_bytes(
            company=config.COMPANY_NAME,
            client_name=shown,
            address=m.get("address") or "",
            month=month,
            lines=lines,
            email=m.get("email") or "",
            intro=letter.get("cover_note"),
            greeting=letter.get("greeting"),
            closing=letter.get("closing"),
            signoff=letter.get("signoff"),
        )
        key = f"bills/{month}/{cid}_{shown.replace(' ', '_')[:40]}.pdf"
        storage.put_bytes(pdf, key, content_type="application/pdf")
        bid = db.upsert_bill(conn, month=month, client_id=cid, s3_key=key)
        out.append({"bill_id": bid, "client_id": cid, "client_name": shown,
                    "email": m.get("email"), "s3_key": key})
    return out


def get_editable_bill(conn, month: str, client_id: int) -> dict | None:
    from . import db

    client = db.get_client(conn, client_id)
    if not client:
        return None
    lines = db.work_for_client_month(conn, month, client_id)
    bill = conn.execute(
        """
        SELECT * FROM bills WHERE month = %s AND client_id = %s
        ORDER BY id DESC LIMIT 1
        """,
        (month, client_id),
    ).fetchone()
    shown_name = client["name"]
    if bill and str(bill.get("display_name") or "").strip():
        shown_name = str(bill.get("display_name")).strip()
    bill_choice = str(bill.get("bill_delivery") or "").strip().lower() if bill else ""
    if bill and bill.get("bill_email") is not None:
        shown_email = bill.get("bill_email") or ""
    else:
        shown_email = client.get("email") or ""
    return {
        "month": month,
        "client_id": client_id,
        "client_name": shown_name,
        "email": shown_email,
        "phone": client.get("phone") or "",
        "address": client.get("address") or "",
        "delivery": bill_choice if bill_choice in {"email", "sms", "mail"} else delivery_channel(client),
        "s3_key": bill["s3_key"] if bill else None,
        "bill_id": bill["id"] if bill else None,
        "greeting": letter_or_default(bill.get("greeting") if bill else None, default_greeting(shown_name)),
        "intro": invoice_sentence(month, bill.get("cover_note") if bill else None),
        "closing": letter_or_default(bill.get("closing") if bill else None, default_closing()),
        "signoff": letter_or_default(bill.get("signoff") if bill else None, default_signoff()),
        "lines": [_editor_line(ln) for ln in lines],
        "preview": [
            {"date": row["date"], "description": row["description"], "amount": float(row["amount"])}
            for row in compile_invoice_lines(lines, month)
        ],
    }


def _editor_line(ln: dict) -> dict:
    role, prior_month = line_role(ln)
    amount = float(ln["amount"]) if ln.get("amount") is not None else 0.0
    return {
        "id": int(ln["id"]),
        "description": ln.get("description") or "",
        "amount": amount,
        "kind": role,
        "prior_month": prior_month,
    }


def _note_for_line(ln: dict, existing: dict | None) -> str | None:
    role = str(ln.get("kind") or "").strip() or line_role(ln)[0]
    prior_month = str(ln.get("prior_month") or "").strip()
    desc = str(ln.get("description") or "").strip()
    if role == "discount" or desc.lower().startswith("discount"):
        return "discount"
    if role == "prior_taxed":
        return f"prior-taxed:{prior_month}"
    if role == "prior":
        return f"prior:{prior_month}"
    if existing and existing.get("day_or_note"):
        return existing.get("day_or_note")
    return None


def save_editable_bill(conn, month: str, client_id: int, *, email: str, address: str,
                       lines: list[dict], intro: str | None = None, phone: str | None = None,
                       delivery: str | None = None, greeting: str | None = None,
                       closing: str | None = None, signoff: str | None = None) -> dict:
    """Update client + work lines, regenerate PDF, return updated edit payload."""
    from . import db

    client = db.get_client(conn, client_id)
    if not client:
        raise ValueError("client not found")
    choice = str(delivery or "").strip().lower()
    if choice not in {"email", "sms", "mail"}:
        choice = delivery_channel(client)
    db.update_client_contact(
        conn,
        client_id,
        email=email or None,
        address=address or None,
        phone=phone or None,
        delivery=choice,
        prefer_mail=choice == "mail",
    )

    stored = {int(r["id"]): r for r in db.work_for_client_month(conn, month, client_id)}
    existing_ids = set(stored)
    keep_ids = set()
    for ln in lines:
        desc = str(ln.get("description") or "").strip()
        raw_amt = ln.get("amount")
        try:
            amt = float(raw_amt) if raw_amt is not None and str(raw_amt).strip() != "" else 0.0
        except (TypeError, ValueError):
            amt = 0.0
        wid = ln.get("id")
        existing = stored.get(int(wid)) if wid and int(wid) in existing_ids else None
        note = _note_for_line(ln, existing)
        if existing:
            db.update_work_item(
                conn, int(wid), description=desc, amount=amt, day_or_note=note,
            )
            keep_ids.add(int(wid))
        elif desc or amt:
            new_id = db.add_work_item(
                conn,
                client_id=client_id,
                month=month,
                day_or_note=note,
                description=desc,
                amount=amt,
            )
            keep_ids.add(new_id)
    for wid in existing_ids - keep_ids:
        db.delete_work_item(conn, wid)

    rows = [dict(r) for r in db.work_for_client_month(conn, month, client_id)]
    client = db.get_client(conn, client_id)
    for row in rows:
        row["mow_price"] = client.get("mow_price")
        row["hedge_price"] = client.get("hedge_price")
    _reprice_visits(conn, rows)
    letter = db.bill_letter(conn, month, client_id)
    shown = str(letter.get("display_name") or "").strip() or client["name"]
    pdf = build_pdf_bytes(
        company=config.COMPANY_NAME,
        client_name=shown,
        address=client.get("address") or "",
        month=month,
        lines=rows,
        email=client.get("email") or "",
        intro=intro,
        greeting=greeting,
        closing=closing,
        signoff=signoff,
    )
    key = f"bills/{month}/{client_id}_{shown.replace(' ', '_')[:40]}.pdf"
    storage.put_bytes(pdf, key, content_type="application/pdf")
    bid = db.upsert_bill(conn, month=month, client_id=client_id, s3_key=key)
    db.set_bill_letter(
        conn,
        bid,
        note=stored_letter(intro, invoice_sentence(month)),
        greeting=stored_letter(greeting, default_greeting(shown)),
        closing=stored_letter(closing, default_closing()),
        signoff=stored_letter(signoff, default_signoff()),
    )
    db.set_bill_face(
        conn,
        bid,
        display_name=str(letter.get("display_name") or "").strip() or None,
        bill_email=None,
        bill_delivery=None,
    )
    return {
        "bill_id": bid,
        "client_id": client_id,
        "client_name": shown,
        "email": client.get("email"),
        "s3_key": key,
        "month": month,
    }


def apply_bill_face(
    conn,
    month: str,
    client_id: int,
    *,
    name: str,
    email: str,
    delivery: str,
    save_to_client: bool,
) -> dict:
    """Rebuild this month's PDF from the list. Keep the client list unless asked."""
    from . import db

    client = db.get_client(conn, client_id)
    if not client:
        raise ValueError("Client was not found")
    bill = conn.execute(
        """
        SELECT * FROM bills WHERE month = %s AND client_id = %s
        ORDER BY id DESC LIMIT 1
        """,
        (month, client_id),
    ).fetchone()
    if not bill:
        raise ValueError("There is no bill for this month yet")
    shown = " ".join(str(name or "").split())
    if not shown:
        raise ValueError("Enter a name")
    choice = str(delivery or "").strip().lower()
    if choice not in {"email", "sms", "mail"}:
        raise ValueError("Delivery must be email, SMS, or paper")
    email_text = str(email or "").strip()
    if save_to_client:
        other = db.get_client_by_name(conn, shown)
        if other and int(other["id"]) != int(client_id):
            raise ValueError(f"{shown} is already on the client list")
        db.update_client_contact(
            conn,
            client_id,
            email=email_text or None,
            address=client.get("address"),
            name=shown,
            phone=client.get("phone"),
            delivery=choice,
            prefer_mail=choice == "mail",
        )
        db.set_bill_face(conn, int(bill["id"]), display_name=None, bill_email=None, bill_delivery=None)
    else:
        roster_name = str(client.get("name") or "").strip()
        roster_email = str(client.get("email") or "").strip()
        roster_delivery = delivery_channel(client)
        db.set_bill_face(
            conn,
            int(bill["id"]),
            display_name=None if shown == roster_name else shown,
            bill_email=None if email_text == roster_email else email_text,
            bill_delivery=None if choice == roster_delivery else choice,
        )
    rows = [dict(row) for row in db.work_for_client_month(conn, month, client_id)]
    priced = db.get_client(conn, client_id) or client
    for row in rows:
        row["mow_price"] = priced.get("mow_price")
        row["hedge_price"] = priced.get("hedge_price")
    letter = db.bill_letter(conn, month, client_id)
    greeting = letter.get("greeting")
    if not _letter_text(greeting) or _letter_text(greeting) == _letter_text(default_greeting(client.get("name") or "")):
        greeting = None
        db.set_bill_letter(
            conn,
            int(bill["id"]),
            note=letter.get("cover_note"),
            greeting=None,
            closing=letter.get("closing"),
            signoff=letter.get("signoff"),
        )
    pdf = build_pdf_bytes(
        company=config.COMPANY_NAME,
        client_name=shown,
        address=priced.get("address") or "",
        month=month,
        lines=rows,
        email=email_text,
        intro=letter.get("cover_note"),
        greeting=greeting,
        closing=letter.get("closing"),
        signoff=letter.get("signoff"),
    )
    key = f"bills/{month}/{client_id}_{shown.replace(' ', '_')[:40]}.pdf"
    storage.put_bytes(pdf, key, content_type="application/pdf")
    bid = db.upsert_bill(conn, month=month, client_id=client_id, s3_key=key)
    return {
        "bill_id": bid,
        "client_id": client_id,
        "client_name": shown,
        "email": email_text,
        "delivery": choice,
        "s3_key": key,
        "month": month,
        "saved_to_client": bool(save_to_client),
    }


def tax_table_tsv(conn, month: str) -> str:
    from . import db

    rows = db.work_for_month(conn, month)
    lines = ["client\tdescription\tamount\tmonth"]
    for r in rows:
        amount = r.get("amount")
        shown = "" if amount is None else amount
        lines.append(
            f"{r['client_name']}\t{r.get('description') or ''}\t{shown}\t{month}"
        )
    return "\n".join(lines) + "\n"


def tax_client_figures(lines: list[dict]) -> dict:
    """This month's revenue, earlier unpaid balances, and the sales tax on the bill.

    Revenue is this month's work after discounts. An earlier unpaid total is not
    revenue. One that has not been taxed yet is included in this month's sales
    tax. One that already includes tax is added after the tax.
    """
    revenue = Decimal("0")
    prior = Decimal("0")
    prior_taxed = Decimal("0")
    prior_notes: list[str] = []
    prior_taxed_notes: list[str] = []
    for line in lines:
        role, prior_month = line_role(line)
        label, _day = _service(line)
        if not label or label.lower() in {"sales tax", "total"}:
            continue
        amount = _money(line.get("amount"))
        if role == "discount":
            revenue += -abs(amount)
        elif role == "prior":
            prior += amount
            if prior_month:
                prior_notes.append(_month_label(prior_month))
        elif role == "prior_taxed":
            prior_taxed += amount
            if prior_month:
                prior_taxed_notes.append(_month_label(prior_month))
        else:
            revenue += amount
    revenue = revenue.quantize(_MONEY, rounding=ROUND_HALF_UP)
    prior = prior.quantize(_MONEY, rounding=ROUND_HALF_UP)
    prior_taxed = prior_taxed.quantize(_MONEY, rounding=ROUND_HALF_UP)
    tax = ((revenue + prior) * _TAX_RATE).quantize(_MONEY, rounding=ROUND_HALF_UP)
    total = (revenue + prior + tax + prior_taxed).quantize(_MONEY, rounding=ROUND_HALF_UP)
    notes = []
    if prior_notes:
        notes.append(", ".join(prior_notes) + " — taxed with this month")
    if prior_taxed_notes:
        notes.append(", ".join(prior_taxed_notes) + " — tax already included")
    return {
        "revenue": revenue,
        "prior": prior,
        "prior_taxed": prior_taxed,
        "tax": tax,
        "total": total,
        "note": "; ".join(notes),
    }


def _month_label(raw: str) -> str:
    text = str(raw or "").strip()
    parts = text.split("-", 1)
    if len(parts) == 2 and parts[0].isdigit() and parts[1].isdigit():
        year, mon = int(parts[0]), int(parts[1])
        if 1 <= mon <= 12:
            return date(year, mon, 1).strftime("%B %Y")
    return text


def tax_table_xlsx(conn, month: str) -> bytes:
    """Excel workbook: this month's revenue, sales tax, and any earlier unpaid balance."""
    from . import db

    rows = db.work_for_month(conn, month)
    by_client: dict[int, list] = defaultdict(list)
    order: list[int] = []
    names: dict[int, str] = {}
    for row in rows:
        cid = int(row["client_id"])
        if cid not in by_client:
            order.append(cid)
            names[cid] = row["client_name"]
        by_client[cid].append(row)
    figures = [(names[cid], tax_client_figures(by_client[cid])) for cid in order]
    return tax_workbook_bytes(month, figures)


def tax_workbook_bytes(month: str, figures: list[tuple[str, dict]]) -> bytes:
    """Build the tax workbook from per-client figures."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Side, Border

    _letter, work_month, _month_num = _invoice_when(month)
    del _letter, _month_num
    year = ""
    parts = str(month or "").split("-", 1)
    if len(parts) == 2 and parts[0].isdigit():
        year = parts[0]
    title = f"Tax table for {work_month} {year}".strip()

    book = Workbook()
    sheet = book.active
    sheet.title = "Tax"
    headers = [
        "Client",
        "This month's revenue",
        "Previous unpaid, taxed this month",
        "Sales tax",
        "Previous unpaid, tax already included",
        "Total billed",
        "Previous unpaid from",
    ]
    sheet.append([title])
    sheet.append([])
    sheet.append(headers)
    header_row = 3
    first_data = 4
    for name, fig in figures:
        sheet.append([
            name,
            float(fig["revenue"]),
            float(fig["prior"]),
            float(fig["tax"]),
            float(fig["prior_taxed"]),
            float(fig["total"]),
            fig["note"],
        ])
    last_data = header_row + len(figures)
    sheet.append([])
    totals = {
        "revenue": sum((fig["revenue"] for _name, fig in figures), Decimal("0")),
        "prior": sum((fig["prior"] for _name, fig in figures), Decimal("0")),
        "tax": sum((fig["tax"] for _name, fig in figures), Decimal("0")),
        "prior_taxed": sum((fig["prior_taxed"] for _name, fig in figures), Decimal("0")),
        "total": sum((fig["total"] for _name, fig in figures), Decimal("0")),
    }
    summary = [
        ("This month's revenue", totals["revenue"]),
        ("Sales tax", totals["tax"]),
        ("Previous unpaid, taxed this month", totals["prior"]),
        ("Previous unpaid, tax already included", totals["prior_taxed"]),
        ("Total billed", totals["total"]),
    ]
    summary_start = last_data + 2
    for label, amount in summary:
        sheet.append([label, float(amount)])
    if totals["prior"] == 0 and totals["prior_taxed"] == 0:
        sheet.append(["No previous unpaid amounts on these bills."])
    else:
        sheet.append(["Previous unpaid amounts are listed above. They are not this month's revenue."])

    money = '"$"#,##0.00'
    thin = Border(
        left=Side(style="thin", color="D0D0D0"),
        right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"),
        bottom=Side(style="thin", color="D0D0D0"),
    )
    header_fill = PatternFill("solid", fgColor="1F4E36")
    tax_fill = PatternFill("solid", fgColor="FFF2CC")
    unpaid_fill = PatternFill("solid", fgColor="FCE4D6")
    title_font = Font(name="Calibri", bold=True, size=14)
    head_font = Font(name="Calibri", bold=True, color="FFFFFF")
    sheet["A1"].font = title_font
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    for col in range(1, len(headers) + 1):
        cell = sheet.cell(header_row, col)
        cell.font = head_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in range(first_data, last_data + 1):
        for col in range(1, len(headers) + 1):
            cell = sheet.cell(row, col)
            cell.border = thin
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for col in (2, 3, 4, 5, 6):
            sheet.cell(row, col).number_format = money
        sheet.cell(row, 4).fill = tax_fill
        if float(sheet.cell(row, 3).value or 0):
            sheet.cell(row, 3).fill = unpaid_fill
            sheet.cell(row, 7).font = Font(name="Calibri", bold=True)
        if float(sheet.cell(row, 5).value or 0):
            sheet.cell(row, 5).fill = unpaid_fill
            sheet.cell(row, 7).font = Font(name="Calibri", bold=True)
    for offset, (label, _amount) in enumerate(summary):
        row = summary_start + offset
        label_cell = sheet.cell(row, 1)
        value_cell = sheet.cell(row, 2)
        label_cell.font = Font(name="Calibri", bold=True)
        value_cell.font = Font(name="Calibri", bold=True)
        value_cell.number_format = money
        if label == "Sales tax":
            label_cell.fill = tax_fill
            value_cell.fill = tax_fill
        if label.startswith("Previous unpaid") and float(value_cell.value or 0):
            label_cell.fill = unpaid_fill
            value_cell.fill = unpaid_fill
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A{header_row}:G{last_data if figures else header_row}"
    widths = [34, 24, 34, 16, 40, 16, 46]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.row_dimensions[header_row].height = 32
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def _has_email(bill: dict) -> bool:
    return delivery_channel(bill) == "email"


def zip_bills(conn, month: str, *, mode: str = "all") -> bytes:
    """Zip PDFs. mode: all | mailing_only (paper) | sms_only | with_email."""
    import zipfile
    from . import db

    bills = db.bills_for_month(conn, month)
    if mode == "mailing_only":
        bills = [b for b in bills if delivery_channel(b) == "mail"]
    elif mode == "sms_only":
        bills = [b for b in bills if delivery_channel(b) == "sms"]
    elif mode == "with_email":
        bills = [b for b in bills if delivery_channel(b) == "email"]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for b in bills:
            data = storage.get_bytes(b["s3_key"])
            name = f"{b['client_name'].replace(' ', '_')}.pdf"
            zf.writestr(name, data)
    return buf.getvalue()
