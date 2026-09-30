"""Money, name, and delivery rules. Uses placeholder clients only."""

from app.billing import _first_name, compile_invoice_lines, contact_email, delivery_channel
from ocr.work_marks import parse_work_marks


def _kinds(text):
    return [(m["kind"], m.get("day"), m.get("name"), m.get("amount")) for m in parse_work_marks(text)]


def test_custom_amount_without_dollar():
    marks = _kinds("bush trimming 50")
    assert ("custom", None, "bush trimming", 50.0) in marks


def test_address_after_price_is_not_a_visit():
    marks = _kinds("trimming $50 12 Main St")
    kinds = [m[0] for m in marks]
    assert kinds == ["custom"]
    assert marks[0][3] == 50.0
    marks = _kinds("50 12 Main St")
    assert [m[0] for m in marks] == ["custom"]
    assert marks[0][3] == 50.0
    marks = _kinds("$50 12 Main St")
    assert [m[0] for m in marks] == ["custom"]


def test_adjacent_numbers_stay_apart():
    marks = _kinds("trimming $50 12")
    assert ("custom",  None, "trimming", 50.0) in [(a, b, c, d) for a, b, c, d in marks] or (
        "custom", None, "trimming", 50.0
    ) in marks
    # 12 follows the price with no street, so it can still be a mow. It must not become 5012.
    assert all(m[3] != 5012 for m in marks)


def test_day_written_with_a_job_is_the_job_date():
    marks = _kinds("5 bush trimming $50")
    assert ("custom", 5, "bush trimming", 50.0) in marks
    assert ("mow", 5, "", None) not in marks
    marks = _kinds("9 16 5 bush trimming 50")
    assert ("mow", 9, "", None) in marks
    assert ("mow", 16, "", None) in marks
    assert ("mow", 5, "", None) not in marks
    assert ("custom", 5, "bush trimming", 50.0) in marks
    marks = _kinds("5 14 21 21 | Hedge Trimming 87 | 29")
    mow_days = [m[1] for m in marks if m[0] == "mow"]
    assert mow_days == [5, 14, 21, 29]
    assert ("custom", 21, "Hedge Trimming", 87.0) in marks
    marks = _kinds("8 | bush trimming $50")
    assert marks[0][0] == "custom" and marks[0][1] == 8
    rows = compile_invoice_lines([
        {"description": "Bush trimming", "amount": 50, "day_or_note": "5"},
        {"description": "Mowing 8", "amount": 0, "day_or_note": "8", "mow_price": 42},
        {"description": "Mowing 15", "amount": 0, "day_or_note": "15", "mow_price": 42},
    ], "2026-09")
    by_name = {row["description"]: row for row in rows}
    assert by_name["Bush trimming"]["date"] == "9/5"
    assert by_name["Mowing"]["amount"] == __import__("decimal").Decimal("84.00")
    assert by_name["Mowing"]["date"] == "9/8, 9/15"


def test_hedge_marks():
    assert _kinds("15h")[0][0] == "hedge"
    assert _kinds("15H")[0][0] == "hedge"
    assert _kinds("15 h")[0][0] == "hedge"
    assert _kinds("15 hedge")[0][0] == "hedge"
    marks = _kinds("9h bush $40")
    assert ("hedge", 9, "", None) in marks
    assert ("custom", None, "bush", 40.0) in marks or ("custom", 9, "bush", 40.0) in marks


def test_tax_includes_pretax_prior_and_not_an_already_taxed_total():
    base = [{"description": "Mowing", "amount": 100, "day_or_note": "5"}]
    rows = compile_invoice_lines(base + [{
        "description": "Previous bill (2026-08)",
        "amount": 50,
        "day_or_note": "prior:2026-08",
    }], "2026-09")
    by_name = {row["description"]: row for row in rows}
    assert by_name["Sales tax"]["amount"] == __import__("decimal").Decimal("9.53")
    assert by_name["Total"]["amount"] == __import__("decimal").Decimal("159.53")
    assert by_name["Previous bill (2026-08)"]["date"] == ""

    rows = compile_invoice_lines(base + [{
        "description": "Previous bill (2026-08)",
        "amount": 50,
        "day_or_note": "prior-taxed:2026-08",
    }], "2026-09")
    by_name = {row["description"]: row for row in rows}
    assert by_name["Sales tax"]["amount"] == __import__("decimal").Decimal("6.35")
    assert by_name["Total"]["amount"] == __import__("decimal").Decimal("156.35")


def test_tax_workbook_separates_revenue_tax_and_previous_unpaid():
    from io import BytesIO
    from openpyxl import load_workbook
    from app.billing import tax_client_figures, tax_workbook_bytes

    taxed_now = tax_client_figures([
        {"description": "Mowing", "amount": 100, "day_or_note": "5"},
        {"description": "Previous bill (2026-08)", "amount": 50, "day_or_note": "prior:2026-08"},
    ])
    assert taxed_now["revenue"] == __import__("decimal").Decimal("100.00")
    assert taxed_now["prior"] == __import__("decimal").Decimal("50.00")
    assert taxed_now["tax"] == __import__("decimal").Decimal("9.53")
    already = tax_client_figures([
        {"description": "Mowing", "amount": 40, "day_or_note": "2"},
        {"description": "Previous bill (2026-08)", "amount": 20, "day_or_note": "prior-taxed:2026-08"},
    ])
    assert already["revenue"] == __import__("decimal").Decimal("40.00")
    assert already["tax"] == __import__("decimal").Decimal("2.54")
    assert already["prior_taxed"] == __import__("decimal").Decimal("20.00")
    plain = tax_client_figures([{"description": "Mowing", "amount": 10, "day_or_note": "1"}])
    data = tax_workbook_bytes("2026-09", [
        ("SMITH, Mary", taxed_now),
        ("JONES, Ann", already),
        ("LEE, Pat", plain),
    ])
    book = load_workbook(BytesIO(data))
    sheet = book.active
    assert sheet["A1"].value.startswith("Tax table for September 2026")
    assert sheet["B3"].value == "This month's revenue"
    assert sheet["D3"].value == "Sales tax"
    assert "Previous unpaid" in sheet["C3"].value
    assert sheet["B4"].value == 100
    assert sheet["C4"].value == 50
    assert sheet["D4"].value == 9.53
    assert "August 2026" in sheet["G4"].value
    assert sheet["E5"].value == 20
    assert sheet["C6"].value == 0
    labels = [sheet.cell(row, 1).value for row in range(8, 14)]
    assert "This month's revenue" in labels
    assert "Sales tax" in labels
    assert any(str(label).startswith("Previous unpaid") for label in labels)


def test_discount_reduces_tax_and_has_no_date():
    rows = compile_invoice_lines([
        {"description": "Mowing", "amount": 100, "day_or_note": "5"},
        {"description": "Discount", "amount": 10, "day_or_note": "discount"},
    ], "2026-09")
    by_name = {row["description"]: row for row in rows}
    assert by_name["Discount"]["date"] == ""
    assert by_name["Discount"]["amount"] == __import__("decimal").Decimal("-10.00")
    assert by_name["Sales tax"]["amount"] == __import__("decimal").Decimal("5.72")


def test_dear_is_the_given_name():
    assert _first_name("SMITH, Mary") == "Mary"
    assert _first_name("AGOSTINI, Lillian") == "Lillian"
    assert _first_name("SMITH, MARY") == "Mary"
    assert _first_name("SMITH Mary") == "Mary"
    assert _first_name("Kelly Small") == "Kelly"


def test_delivery():
    assert delivery_channel({"email": "a@b.com", "phone": "860"}) == "email"
    assert delivery_channel({"email": "", "billing_notes": "REGULAR MAIL. a@b.com", "phone": "860"}) == "email"
    assert contact_email({"email": "", "billing_notes": "see a@b.com"}) == "a@b.com"
    assert delivery_channel({"email": "", "phone": "860", "billing_notes": ""}) == "sms"
    assert delivery_channel({"email": "", "phone": "", "billing_notes": ""}) == "mail"
    assert delivery_channel({"email": "a@b.com", "phone": "860", "prefer_mail": True}) == "mail"
    assert delivery_channel({"email": "", "phone": "860", "prefer_mail": True}) == "mail"
    assert delivery_channel({"delivery": "sms", "email": "a@b.com", "phone": "860"}) == "sms"
    assert delivery_channel({"delivery": "mail", "email": "a@b.com"}) == "mail"
    assert delivery_channel({"delivery": "email", "phone": "860"}) == "email"


def test_bill_face_uses_this_months_name_email_and_delivery():
    from app.billing import delivery_channel, face_bill

    row = face_bill({
        "client_name": "SMITH, Mary",
        "email": "mary@example.com",
        "delivery": "email",
        "display_name": "SMITH, Marie",
        "bill_email": "",
        "bill_delivery": "mail",
        "prefer_mail": False,
    })
    assert row["client_name"] == "SMITH, Marie"
    assert row["email"] == ""
    assert delivery_channel(row) == "mail"
    plain = face_bill({
        "client_name": "SMITH, Mary",
        "email": "mary@example.com",
        "phone": "",
        "prefer_mail": False,
        "bill_email": None,
        "bill_delivery": None,
    })
    assert plain["client_name"] == "SMITH, Mary"
    assert delivery_channel(plain) == "email"


def test_repeated_work_chunk_is_one_review_row():
    from ocr.guided import _assemble_work_rows

    columns = ["CLIENT", "DATE & WORK COMPLETED"]
    names = ["SMITH, Mary", "JONES, Ann", "LEE, Pat", "DIAZ, Omar", "NGUYEN, Kim"]
    once = [
        ["SMITH, Mary", "9 16"],
        ["JONES, Ann", "15h"],
        ["LEE, Pat", "5 bush trimming 50"],
        ["DIAZ, Omar", "12"],
        ["NGUYEN, Kim", "3 10"],
    ]
    # Two chunks each returned the whole sheet.
    rows = _assemble_work_rows(names, once + once, columns)
    assert [row[0] for row in rows] == names
    assert [row[1] for row in rows] == [row[1] for row in once]


def test_richer_copy_wins_and_unmatched_name_stays_once():
    from ocr.guided import _assemble_work_rows

    columns = ["CLIENT", "DATE & WORK COMPLETED"]
    names = ["SMITH, Mary"]
    pool = [
        ["SMITH, Mary", "9"],
        ["SMITH, Mary", "9 16 23"],
        ["ADDED, Bob", "2"],
        ["ADDED, Bob", "2"],
    ]
    rows = _assemble_work_rows(names, pool, columns)
    assert rows == [
        ["SMITH, Mary", "9 16 23"],
        ["ADDED, Bob", "2"],
    ]


def test_invoice_sentence_can_be_replaced():
    from app.billing import invoice_sentence

    assert invoice_sentence("2026-09") == "Below is the invoice for any work done in September."
    custom = "Below is the invoice for work done in September, plus the August balance."
    assert invoice_sentence("2026-09", custom) == custom
    assert invoice_sentence("2026-09", "  ") == "Below is the invoice for any work done in September."


def test_letter_wording_restores_when_cleared():
    from app.billing import default_closing, default_greeting, default_signoff, letter_or_default, stored_letter

    assert default_greeting("SMITH, Mary") == "Dear Mary,"
    assert "checks payable" in default_closing()
    assert default_signoff().startswith("With regards,")
    assert letter_or_default("  ", default_greeting("SMITH, Mary")) == "Dear Mary,"
    assert letter_or_default("Hello Mary,", default_greeting("SMITH, Mary")) == "Hello Mary,"
    assert stored_letter("Dear Mary,", default_greeting("SMITH, Mary")) is None
    assert stored_letter("Hello Mary,", default_greeting("SMITH, Mary")) == "Hello Mary,"


def test_unknown_name_is_a_conflict_until_answered():
    from app.knowledge import open_conflicts

    extract = {
        "tables": [{
            "columns": ["CLIENT", "DATE & WORK COMPLETED"],
            "rows": [
                ["SMITH, Mary", "9 16"],
                ["ZZZ, Test", "15h"],
                ["ZZZ, Test", "bush 40"],
            ],
        }],
    }
    conflicts = open_conflicts(extract, ["SMITH Mary"], [])
    assert [row["name"] for row in conflicts] == ["ZZZ, Test"]
    assert "15h" in conflicts[0]["work"]
    assert "bush 40" in conflicts[0]["work"]
    closed = open_conflicts(extract, ["SMITH, Mary"], [{"name": "ZZZ, Test", "add_permanently": True}])
    assert closed == []


def test_near_spelling_suggests_the_roster_client():
    from app.knowledge import open_conflicts

    extract = {
        "tables": [{
            "columns": ["CLIENT", "DATE & WORK COMPLETED"],
            "rows": [["SMIHT, Mary", "9 16"]],
        }],
    }
    roster = [
        {"id": 4, "name": "SMITH, Mary"},
        {"id": 5, "name": "JONES, Ann"},
    ]
    conflicts = open_conflicts(extract, roster, [])
    assert conflicts[0]["suggestion"]["id"] == 4
    assert conflicts[0]["suggestion"]["name"] == "SMITH, Mary"
    close = {
        "tables": [{
            "columns": ["CLIENT", "DATE & WORK COMPLETED"],
            "rows": [["SMITH, Marie", "4"]],
        }],
    }
    assert open_conflicts(close, roster, [])[0]["suggestion"]["name"] == "SMITH, Mary"
    other = {
        "tables": [{
            "columns": ["CLIENT", "DATE & WORK COMPLETED"],
            "rows": [["SMITH, John", "2"]],
        }],
    }
    assert open_conflicts(other, roster, [])[0]["suggestion"] is None


def test_set_up_client_is_used_for_a_near_spelling():
    from app.knowledge import filing_client, open_conflicts

    roster = [
        {"id": 4, "name": "SMITH, Mary", "address": "123 Main St.", "email": "mary@example.com", "mow_price": 40},
        {"id": 8, "name": "SMIHT, Mary"},
        {"id": 9, "name": "MMM, Test", "address": "9 Oak", "phone": "8605550100"},
    ]
    assert filing_client("SMIHT, Mary", roster)["id"] == 4
    assert filing_client("Test MMM", roster)["id"] == 9
    assert filing_client("SMITH, John", roster) is None
    near = {
        "tables": [{
            "columns": ["CLIENT", "DATE & WORK COMPLETED"],
            "rows": [["SMIHT, Mary", "9"]],
        }],
    }
    assert open_conflicts(near, roster, []) == []
    other = {
        "tables": [{
            "columns": ["CLIENT", "DATE & WORK COMPLETED"],
            "rows": [["SMITH, John", "2"]],
        }],
    }
    assert open_conflicts(other, roster, [])[0]["name"] == "SMITH, John"


def test_empty_work_rows_are_left_out():
    from ocr.guided import _assemble_work_rows
    from ocr.prompts import guided_work_prompt

    columns = ["CLIENT", "DATE & WORK COMPLETED"]
    names = ["SMITH, Mary", "BLANK, Pat", "JONES, Ann"]
    pool = [
        ["SMITH, Mary", "9 16"],
        ["BLANK, Pat", ""],
        ["JONES, Ann", "15h"],
    ]
    rows = _assemble_work_rows(names, pool, columns)
    assert [row[0] for row in rows] == ["SMITH, Mary", "JONES, Ann"]
    prompt = guided_work_prompt(["BLANK, Pat"], columns)
    assert "Skip a row when the work cell is blank" in prompt
    assert "If the work cell is blank on the paper, use" not in prompt

